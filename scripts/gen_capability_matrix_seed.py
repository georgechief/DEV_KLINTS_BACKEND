"""Generate dataruns/capabilities/matrix_seed.json from pack Matrix xlsx sheet 02.

Usage (from klints_backend):
  python scripts/gen_capability_matrix_seed.py

Runtime does not need the xlsx — only the committed JSON seed.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
XLSX_REL = (
    "Klints_MVP1_Rohan_Build_Pack_v1.2_20260718"
    "/02_Execution_Capabilities"
    "/Klints_Manago_ExecutionCapabilityMatrix_v1.1_20260718.xlsx"
)
OUT_REL = "dataruns/capabilities/matrix_seed.json"
SHEET_NAME = "02 Capability Matrix"
PACK_OBSERVED_AT = "2026-07-18T12:00:00+02:00"
SOURCE_LABEL = "pack:Capability Matrix v1.1"

# Sheet 05 Workflow build: UPSERT fallback is HUMAN.WORKFLOW.BUILD (capability id).
FALLBACK_ID_OVERRIDES = {
    "MCP.WORKFLOW.UPSERT": "HUMAN.WORKFLOW.BUILD",
}

MUST_SEED_STATUS = {
    "HUMAN.WORKFLOW.BUILD": "CONFIRMED_LIVE",
    "MCP.WORKFLOW.UPSERT": "DISCOVERY_REQUIRED",
    "MCP.WORKFLOW.PUBLISH": "DISCOVERY_REQUIRED",
    "RESTV2.WORKFLOW.LIST": "CONFIRMED_LIVE",
}


def _cell(row: tuple, idx: int) -> str:
    if idx >= len(row) or row[idx] is None:
        return ""
    return str(row[idx]).strip()


def _build_evidence(*, capability_id: str, evidence_source: str, status: str) -> list[dict]:
    """MCP / discovery rows stay empty; confirmed rows cite Matrix evidence source."""
    if status == "DISCOVERY_REQUIRED" or capability_id.startswith("MCP."):
        return []
    if not evidence_source:
        return []
    return [
        {
            "source": SOURCE_LABEL,
            "locator": f"{SHEET_NAME}/{capability_id}",
            "value": evidence_source,
            "observed_at": PACK_OBSERVED_AT,
        }
    ]


def generate_seed(xlsx_path: Path) -> dict:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise SystemExit("openpyxl is required to regenerate the seed") from exc

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        wb.close()
        raise SystemExit(f"Sheet {SHEET_NAME!r} not found in {xlsx_path}")
    ws = wb[SHEET_NAME]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    hdr_i = next(
        (i for i, r in enumerate(rows) if r and str(r[0] or "").strip() == "Capability ID"),
        None,
    )
    if hdr_i is None:
        raise SystemExit("Capability ID header not found")

    capabilities: list[dict] = []
    seen: set[str] = set()
    for row in rows[hdr_i + 1 :]:
        if not row:
            continue
        capability_id = _cell(row, 0)
        if not capability_id:
            continue
        if capability_id in seen:
            raise SystemExit(f"Duplicate capability_id: {capability_id}")
        seen.add(capability_id)

        channel = _cell(row, 3)
        mode = _cell(row, 4)
        status = _cell(row, 5)
        input_contract = _cell(row, 6)
        output_contract = _cell(row, 7)
        fallback_raw = _cell(row, 9)
        evidence_source = _cell(row, 10)

        if capability_id in FALLBACK_ID_OVERRIDES:
            fallback: str | None = FALLBACK_ID_OVERRIDES[capability_id]
        elif fallback_raw:
            fallback = fallback_raw
        else:
            fallback = None

        evidence = _build_evidence(
            capability_id=capability_id,
            evidence_source=evidence_source,
            status=status,
        )
        record = {
            "schema_version": "1.0.0",
            "capability_id": capability_id,
            "channel": channel,
            "mode": mode,
            "status": status,
            "input_contract": input_contract or "",
            "output_contract": output_contract or "",
            "fallback": fallback,
            "evidence": evidence,
            "verified_at": PACK_OBSERVED_AT if evidence else None,
        }
        capabilities.append(record)

    # Validate must-seed
    by_id = {r["capability_id"]: r for r in capabilities}
    for cap_id, expected_status in MUST_SEED_STATUS.items():
        if cap_id not in by_id:
            raise SystemExit(f"Must-seed missing: {cap_id}")
        if by_id[cap_id]["status"] != expected_status:
            raise SystemExit(
                f"Must-seed status mismatch for {cap_id}: "
                f"{by_id[cap_id]['status']!r} != {expected_status!r}"
            )
        if cap_id.startswith("MCP.") and by_id[cap_id]["evidence"]:
            raise SystemExit(f"MCP seed must have empty evidence: {cap_id}")

    return {
        "schema_version": "1.0.0",
        "source": SOURCE_LABEL,
        "generated_from": XLSX_REL.replace("\\", "/") + f"#{SHEET_NAME}",
        "count": len(capabilities),
        "capabilities": capabilities,
    }


def main() -> int:
    xlsx_path = REPO_ROOT / XLSX_REL
    out_path = REPO_ROOT / OUT_REL
    if not xlsx_path.exists():
        print(f"xlsx not found: {xlsx_path}", file=sys.stderr)
        return 1
    seed = generate_seed(xlsx_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(
        json.dumps(seed, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"Wrote {out_path} ({seed['count']} capabilities)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
