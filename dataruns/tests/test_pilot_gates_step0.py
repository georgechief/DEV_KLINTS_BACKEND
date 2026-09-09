"""DCS-09 Step 0 — preconditions (pack inventory + code gap baseline)."""

from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

from django.test import SimpleTestCase

from dataruns.dcs.master import load_check_master_from_json
from dataruns.use_cases.constants import SUPPLEMENTAL_PREFLIGHT_CHECKS

_REPO_ROOT = Path(__file__).resolve().parents[2]
_PACK = _REPO_ROOT / "Klints_MVP1_Rohan_Build_Pack_v1.2_20260718"
_DCS_XLSX = (
    _PACK
    / "01_Specifications"
    / "Klints_Spec_InitialDataConsistencyCheck_v1.4.1_20260718.xlsx"
)
_MANIFEST = _PACK / "04_MVP1_Pilot_Blueprints" / "pilot_manifest.json"
_BLUEPRINTS_DIR = _PACK / "04_MVP1_Pilot_Blueprints"
_ARCHIVE = (
    _REPO_ROOT
    / "docs"
    / "dcs_scoring"
    / "PRD_DCS_09_PILOT_SUPPLEMENTAL_GATES-FUTURE-PRD.md"
)
_SAHIL_PRD = (
    _REPO_ROOT / "docs" / "sahil" / "PRD_DCS_09_PILOT_SUPPLEMENTAL_GATES.md"
)
_WORKING_GAPS = _REPO_ROOT / "docs" / "sahil" / "DCS_09_WORKING_GAPS.md"

# PRD §3.1 + sheet 11 (locked for Step 1 mapping)
_PRD_UC_SUPPLEMENTAL: dict[str, list[str]] = {
    "UC-02": ["CI-08", "CC-06"],
    "UC-04": ["PT-13"],
    "UC-05": ["CC-06"],
    "UC-06B": ["SP-10"],
    "UC-08": ["BR-09"],
    "UC-09": ["LE-10"],
    "UC-11": ["BR-09", "SP-04"],
    "UC-12": ["PT-06"],
    "UC-13": ["PT-05"],
    "UC-21": ["LE-07", "PT-06"],
    "UC-28": ["BR-03", "PT-11"],
}

# MVP1 seeded pilots with no supplemental in blueprint gates (verified Step 0)
_PILOTS_WITHOUT_SUPPLEMENTAL = frozenset(
    {"UC-10", "UC-16", "UC-17", "UC-23", "UC-36"}
)

_SHEET_11 = "11 Pilot Supplemental Gates"
_SHEET_09 = "09 MVP1 Check Scope"
_SHEET_02 = "02 Check Catalogue"


def _blueprint_supplemental_map() -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for path in sorted(_BLUEPRINTS_DIR.glob("UC-*_blueprint.json")):
        body = json.loads(path.read_text(encoding="utf-8"))
        uc = str(body.get("use_case_id") or "").strip()
        gates = body.get("gates") if isinstance(body.get("gates"), dict) else {}
        ids = gates.get("gating_check_ids") if isinstance(gates, dict) else []
        supplemental = sorted(
            str(g).strip()
            for g in (ids or [])
            if str(g).strip() in SUPPLEMENTAL_PREFLIGHT_CHECKS
        )
        if supplemental:
            out[uc] = supplemental
    return out


def _all_blueprint_use_case_ids() -> set[str]:
    out: set[str] = set()
    for path in _BLUEPRINTS_DIR.glob("UC-*_blueprint.json"):
        body = json.loads(path.read_text(encoding="utf-8"))
        uc = str(body.get("use_case_id") or "").strip()
        if uc:
            out.add(uc)
    return out


def _sheet11_uc_map() -> dict[str, list[str]]:
    """Parse sheet 11 Required By → UC → sorted supplemental check ids."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl required for Step 0 sheet 11 inventory") from exc

    wb = openpyxl.load_workbook(_DCS_XLSX, read_only=True, data_only=True)
    try:
        if _SHEET_11 not in wb.sheetnames:
            raise AssertionError(f"missing sheet {_SHEET_11!r}; got {wb.sheetnames}")
        ws = wb[_SHEET_11]
        check_to_ucs: dict[str, list[str]] = {}
        for row in ws.iter_rows(min_row=6, max_row=17, values_only=True):
            if not row or not row[0]:
                continue
            cid = str(row[0]).strip()
            required = str(row[6] or "")
            ucs = [part.strip() for part in required.split(",") if part.strip()]
            check_to_ucs[cid] = ucs
        inverted: dict[str, list[str]] = defaultdict(list)
        for cid, ucs in check_to_ucs.items():
            for uc in ucs:
                inverted[uc].append(cid)
        return {uc: sorted(ids) for uc, ids in inverted.items()}
    finally:
        wb.close()


def _sheet11_check_ids() -> list[str]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl required for Step 0 sheet 11 inventory") from exc

    wb = openpyxl.load_workbook(_DCS_XLSX, read_only=True, data_only=True)
    try:
        ws = wb[_SHEET_11]
        ids: list[str] = []
        for row in ws.iter_rows(min_row=6, max_row=17, values_only=True):
            if row and row[0]:
                ids.append(str(row[0]).strip())
        return ids
    finally:
        wb.close()


class PilotGatesStep0PreconditionsTests(SimpleTestCase):
    """Pack SoT present; 12 IDs locked; zero overlap with headline 42."""

    def test_pack_dcs_workbook_present(self):
        self.assertTrue(_DCS_XLSX.is_file(), msg=str(_DCS_XLSX))
        self.assertGreater(_DCS_XLSX.stat().st_size, 0)

    def test_workbook_has_required_sheets(self):
        import openpyxl

        wb = openpyxl.load_workbook(_DCS_XLSX, read_only=True, data_only=True)
        try:
            names = set(wb.sheetnames)
            self.assertIn(_SHEET_02, names)
            self.assertIn(_SHEET_09, names)
            self.assertIn(_SHEET_11, names)
        finally:
            wb.close()

    def test_sheet11_twelve_ids_match_constants(self):
        ids = _sheet11_check_ids()
        self.assertEqual(len(ids), 12)
        self.assertEqual(set(ids), set(SUPPLEMENTAL_PREFLIGHT_CHECKS))

    def test_sheet11_uc_map_matches_prd(self):
        from_sheet = _sheet11_uc_map()
        expected = {
            uc: sorted(gates) for uc, gates in _PRD_UC_SUPPLEMENTAL.items()
        }
        self.assertEqual(from_sheet, expected)

    def test_pilot_manifest_supplemental_list_matches_constants(self):
        self.assertTrue(_MANIFEST.is_file())
        manifest = json.loads(_MANIFEST.read_text(encoding="utf-8"))
        listed = manifest.get("supplemental_preflight_checks")
        self.assertIsInstance(listed, list)
        self.assertEqual(len(listed), 12)
        self.assertEqual(set(listed), set(SUPPLEMENTAL_PREFLIGHT_CHECKS))

    def test_supplemental_ids_do_not_overlap_headline_42(self):
        master = load_check_master_from_json()
        headline_ids = {c.check_id for c in master.checks}
        self.assertEqual(len(headline_ids), 42)
        overlap = headline_ids & set(SUPPLEMENTAL_PREFLIGHT_CHECKS)
        self.assertEqual(overlap, set(), msg=f"overlap with 42: {sorted(overlap)}")

    def test_uc02_hard_gate_in_42_supplementals_outside_42(self):
        master = load_check_master_from_json()
        headline_ids = {c.check_id for c in master.checks}
        self.assertIn("CC-03", headline_ids)
        self.assertNotIn("CC-06", headline_ids)
        self.assertNotIn("CI-08", headline_ids)

    def test_sixteen_blueprints_present(self):
        paths = list(_BLUEPRINTS_DIR.glob("UC-*_blueprint.json"))
        self.assertEqual(len(paths), 16)
        self.assertEqual(len(_all_blueprint_use_case_ids()), 16)

    def test_blueprint_uc_map_matches_prd_sheet_11(self):
        from_blueprints = _blueprint_supplemental_map()
        self.assertEqual(
            {uc: sorted(gates) for uc, gates in from_blueprints.items()},
            {uc: sorted(gates) for uc, gates in _PRD_UC_SUPPLEMENTAL.items()},
        )

    def test_pilots_without_supplemental_are_expected_set(self):
        all_ucs = _all_blueprint_use_case_ids()
        with_sup = set(_blueprint_supplemental_map())
        without = all_ucs - with_sup
        self.assertEqual(without, _PILOTS_WITHOUT_SUPPLEMENTAL)
        self.assertEqual(len(with_sup) + len(without), 16)

    def test_evaluate_orchestration_exists_after_step5(self):
        """Step 5 delivered evaluate.py; keep asserting the package path."""
        package = _REPO_ROOT / "dataruns" / "dcs" / "pilot_gates"
        self.assertTrue(
            (package / "evaluate.py").is_file(),
            msg="evaluate.py required after Step 5",
        )

    def test_sahil_prd_and_working_gaps_present(self):
        self.assertTrue(_SAHIL_PRD.is_file(), msg=str(_SAHIL_PRD))
        self.assertTrue(_WORKING_GAPS.is_file(), msg=str(_WORKING_GAPS))
        self.assertTrue(_ARCHIVE.is_file(), msg=str(_ARCHIVE))
        gaps = _WORKING_GAPS.read_text(encoding="utf-8")
        self.assertIn("| 0 Preconditions | **Done**", gaps)
